"""Attachment parsers: turn .txt / .pdf / .docx / .xlsx bytes into plain lines.

Every parser returns a ParsedDoc. A document that cannot be read is never an
exception for the caller — it comes back with `readable=False` and an `issue`
string so the pipeline can escalate it to a human with the reason.
"""
from __future__ import annotations

import io
import threading
from dataclasses import dataclass, field

# pdfium (used by pdfplumber for page rendering) is not thread-safe; the batch
# runner processes emails in parallel, so all PDF work is serialised here.
_PDF_LOCK = threading.Lock()


@dataclass
class ParsedDoc:
    path: str
    fmt: str
    lines: list[str] = field(default_factory=list)
    readable: bool = True
    issue: str | None = None          # "corrupt" | "image_only" | "empty" | "unsupported"
    image_pages: list[bytes] = field(default_factory=list)   # PNGs of scanned pages (for vision/OCR)

    @property
    def text(self) -> str:
        return "\n".join(self.lines)


def _fmt(path: str) -> str:
    return path.rsplit(".", 1)[-1].lower() if "." in path else ""


def parse_attachment(path: str, data: bytes) -> ParsedDoc:
    fmt = _fmt(path)
    try:
        if fmt == "txt":
            doc = _parse_txt(path, data)
        elif fmt == "pdf":
            doc = _parse_pdf(path, data)
        elif fmt == "docx":
            doc = _parse_docx(path, data)
        elif fmt in ("xlsx", "xlsm"):
            doc = _parse_xlsx(path, data)
        elif fmt in ("png", "jpg", "jpeg", "tif", "tiff", "bmp", "webp"):
            doc = _parse_image(path, data)
        else:
            return ParsedDoc(path, fmt, readable=False, issue="unsupported")
    except Exception as exc:  # corrupt / truncated / password-protected file
        return ParsedDoc(path, fmt, readable=False, issue=f"corrupt ({type(exc).__name__})")
    if doc.readable and not any(l.strip() for l in doc.lines):
        doc.readable = False
        doc.issue = doc.issue or "empty"
    return doc


def _parse_txt(path: str, data: bytes) -> ParsedDoc:
    text = data.decode("utf-8", errors="replace")
    if text.count("�") > max(5, len(text) // 20):
        return ParsedDoc(path, "txt", readable=False, issue="corrupt (binary content)")
    return ParsedDoc(path, "txt", lines=text.splitlines())


def _parse_pdf(path: str, data: bytes) -> ParsedDoc:
    with _PDF_LOCK:
        return _parse_pdf_unlocked(path, data)


def _parse_pdf_unlocked(path: str, data: bytes) -> ParsedDoc:
    import pdfplumber

    lines: list[str] = []
    images: list[bytes] = []
    has_images = False
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        if not pdf.pages:
            return ParsedDoc(path, "pdf", readable=False, issue="corrupt (no pages)")
        for page in pdf.pages:
            txt = page.extract_text() or ""
            if txt.strip():
                lines.extend(_pdf_lines(page) or txt.splitlines())
            elif page.images:
                has_images = True
                try:
                    buf = io.BytesIO()
                    page.to_image(resolution=300).original.convert("L").save(buf, format="PNG")
                    images.append(buf.getvalue())
                except Exception:
                    pass  # still an image-only page; just no preview for the vision model
    if not lines and has_images:
        return ParsedDoc(path, "pdf", readable=False, issue="image_only", image_pages=images)
    return ParsedDoc(path, "pdf", lines=lines)


def _pdf_lines(page) -> list[str]:
    """Rebuild lines from characters so that a long bold label which overflows
    into the value column ("Notify Party/Intermediate Consignee" over "CERIEX")
    does not get interleaved with the value. Label = leading bold run, value =
    the regular-weight run after it; everything else keeps plain reading order."""
    rows: list[list[dict]] = []
    for ch in page.chars:
        if rows and abs(rows[-1][0]["top"] - ch["top"]) < 2:
            rows[-1].append(ch)
        else:
            rows.append([ch])
    rows.sort(key=lambda r: r[0]["top"])

    def bold(c):
        return "bold" in (c.get("fontname") or "").lower()

    out: list[str] = []
    for row in rows:
        flags = [bold(c) for c in row if c["text"].strip()]
        if not flags:
            continue
        n_bold = next((i for i, b in enumerate(flags) if not b), len(flags))
        is_split = 0 < n_bold < len(flags) and not any(flags[n_bold:])
        if is_split:
            label = "".join(c["text"] for c in row if bold(c)).strip().rstrip(":：").strip()
            value = "".join(c["text"] for c in row if not bold(c)).strip()
            out.append(f"{label}: {value}")
        else:
            indent = "  " if row[0]["x0"] > 120 and not any(flags) else ""
            out.append(indent + _join_chars(sorted(row, key=lambda c: c["x0"])))
    return out


def _join_chars(chars: list[dict]) -> str:
    s, prev = "", None
    for c in chars:
        if prev is not None and c["x0"] - prev["x1"] > 1.5 and not s.endswith(" "):
            s += " "
        s += c["text"]
        prev = c
    return s.strip()


def _parse_image(path: str, data: bytes) -> ParsedDoc:
    from PIL import Image

    img = Image.open(io.BytesIO(data))
    buf = io.BytesIO()
    img.convert("L").save(buf, format="PNG")
    return ParsedDoc(path, "image", readable=False, issue="image_only", image_pages=[buf.getvalue()])


def _parse_docx(path: str, data: bytes) -> ParsedDoc:
    import docx

    d = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    # Keep body order: paragraphs and tables interleaved.
    for block in d.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            t = "".join(n.text or "" for n in block.iter() if n.tag.endswith("}t"))
            if t.strip():
                lines.append(t)
        elif tag == "tbl":
            table = docx.table.Table(block, d)
            for row in table.rows:
                cells = _dedupe([c.text.strip() for c in row.cells])
                lines.extend(_row_to_lines(cells))
    return ParsedDoc(path, "docx", lines=lines)


def _parse_xlsx(path: str, data: bytes) -> ParsedDoc:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    lines: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [_cell(v) for v in row]
            cells = [c for c in cells if c]
            if cells:
                lines.extend(_row_to_lines(cells))
    return ParsedDoc(path, "xlsx", lines=lines)


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _dedupe(cells: list[str]) -> list[str]:
    """Merged table cells repeat their text; keep one copy."""
    out: list[str] = []
    for c in cells:
        if c and (not out or out[-1] != c):
            out.append(c)
    return out


def _row_to_lines(cells: list[str]) -> list[str]:
    """A 2-cell row is `label | value`; render it as `label: value` so the same
    field matcher works for tables and plain text. Multi-line cell values become
    indented continuation lines, like the .txt format."""
    if len(cells) == 2:
        label, value = cells
        value = value.replace(" | ", "\n")
        first, *rest = value.splitlines() or [""]
        return [f"{label}: {first}"] + [f"  {r}" for r in rest]
    return [" ".join(cells)]
