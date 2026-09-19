"""Build demo/ — one folder per capability, each with the attachments to upload
and the email text to paste into "New email". Regenerate any time:

    python scripts/make_demo_kit.py
"""
import io
from pathlib import Path

import docx
import openpyxl
from PIL import Image, ImageDraw, ImageFilter, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "demo"

SHIPPER = "APRIL FAR EAST (M) SDN BHD\n  TOWER 2, AVENUE 5, LEVEL 6; BANGSAR SOUTH CITY; 59200 KUALA LUMPUR, MALAYSIA"
CONSIGNEE = "KTP CO., LTD\n  KTP BLDG., 36 SANGWON-GIL; SEONGDONG-GU, SEOUL, SOUTH KOREA"
NOTIFY = "SAME AS CONSIGNEE"


def txt(path, title, rows):
    lines = [title, "=" * 40, ""]
    for k, v in rows:
        first, *rest = v.split("\n")
        lines.append(f"{k}: {first}")
        lines += rest
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def xlsx(path, title, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S.I."
    ws.append(["APRIL FAR EAST (M) SDN BHD"])
    ws.append([])
    ws.append([title, "3658209911"])
    for k, v in rows:
        ws.append([k, v.replace("\n  ", " | ")])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 70
    wb.save(path)


def docx_table(path, title, rows):
    d = docx.Document()
    d.add_heading(title, level=1)
    d.add_paragraph("B/L NO.: KTPSIN2026091   FREIGHT PREPAID")
    t = d.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in rows:
        c = t.add_row().cells
        c[0].text, c[1].text = k, v.replace("\n  ", "\n")
    d.save(path)


def scanned_pdf(path, title, rows):
    """Image-only PDF: rendered text, slightly rotated and blurred like a scan."""
    W, H = 1240, 1754
    img = Image.new("L", (W, H), 250)
    dr = ImageDraw.Draw(img)
    try:
        big, small = ImageFont.truetype("arial.ttf", 38), ImageFont.truetype("arial.ttf", 27)
    except OSError:
        big = small = ImageFont.load_default()
    dr.text((110, 120), title, font=big, fill=20)
    y = 230
    for k, v in rows:
        for i, line in enumerate(v.split("\n")):
            dr.text((110, y), (k + ":" if i == 0 else ""), font=small, fill=35)
            dr.text((470, y), line.strip(), font=small, fill=35)
            y += 44
        y += 14
    img = img.rotate(0.8, fillcolor=245).filter(ImageFilter.GaussianBlur(0.9))
    img.convert("RGB").save(path, "PDF", resolution=150)


def email(path, subject, body):
    path.write_text(f"SUBJECT\n{subject}\n\nMESSAGE\n{body}\n", encoding="utf-8")


def main():
    OUT.mkdir(exist_ok=True)

    # 1 ── messy formats, identical meaning → must NOT raise false alarms
    f = OUT / "1_messy_but_matching"; f.mkdir(exist_ok=True)
    xlsx(f / "SI_KTP.xlsx", "BL INSTRUCTION", [
        ("Shipper/Exporter", SHIPPER), ("Consignee (Non-Negotiable)", "KTP Co., Ltd.\n  36 SANGWON-GIL, SEOUL"),
        ("NOTIFY PARTY", NOTIFY), ("Load Port", "Port Klang (Westport), Malaysia"),
        ("POD", "Ho Chi Minh City, Viet Nam"), ("Container Count", "2 x 40'HC + 1 x 20'GP"),
        ("GROSS WEIGHT", "64.250,00 KGS"), ("NET WEIGHT", "62,900 KGS"), ("Commodity", "PAPERONE DIGITAL COPIER PAPER")])
    docx_table(f / "BL_KTP.docx", "BILL OF LADING (DRAFT)", [
        ("Shipper (Principal or Seller) (发货人)", SHIPPER), ("Consignee (收货人)", CONSIGNEE),
        ("Notify Party (通知人)", NOTIFY), ("PORT OF LOADING (装货港)", "PORT KLANG (WESTPORT), MALAYSIA (MYPKG)"),
        ("Port of Discharge (卸货港)", "HOCHIMINH CITY, VIETNAM (VNSGN)"), ("Total Containers (箱数)", "3 (THREE) CONTAINERS"),
        ("Gross Wt (kgs) (毛重 KGS)", "64.25 MT"), ("Commodity (货名)", "PAPERONE DIGITAL COPIER PAPER")])
    email(f / "email.txt", "RE_ TO CONFIRM DOCS _ 5RSG-90001 _ HOCHIMINH CITY _ KTP CO., LTD",
          "Dear Team,\n\nPlease check the draft BL against the SI and revert with any discrepancy.\n\nThanks,\nMei Ling")

    # 2 ── real discrepancies hidden in similar-looking documents
    f = OUT / "2_real_discrepancies"; f.mkdir(exist_ok=True)
    rows = [("Shipper", SHIPPER), ("Consignee", CONSIGNEE), ("Notify Party", NOTIFY),
            ("Port of Loading", "SINGAPORE (SGSIN)"), ("Port of Discharge", "BUSAN, SOUTH KOREA (KRPUS)"),
            ("No. of Containers", "4 x 40'HC"), ("Gross Weight (KG)", "88,420 KG")]
    txt(f / "SI_busan.txt", "SHIPPING INSTRUCTION", rows)
    bad = dict(rows)
    bad.update({"Port of Discharge": "PUSAN, SOUTH KOREA (KRPUS)", "No. of Containers": "5 x 40'HC", "Gross Weight (KG)": "88,920 KG",
                "Consignee": "KTP CORPORATION\n  KTP BLDG., 36 SANGWON-GIL; SEONGDONG-GU, SEOUL, SOUTH KOREA"})
    docx_table(f / "BL_busan.docx", "BILL OF LADING (DRAFT)", [
        ("SHIPPER", bad["Shipper"]), ("CONSIGNEE", bad["Consignee"]), ("Notify", bad["Notify Party"]),
        ("Load Port", bad["Port of Loading"]), ("POD", bad["Port of Discharge"]),
        ("Total Containers", bad["No. of Containers"]), ("Gross Wt (kgs)", bad["Gross Weight (KG)"])])
    email(f / "email.txt", "Draft BL SOLID 16 V.044NW2 SINGAPORE - amend BL 061",
          "Hi Team,\n\nAttached are the SI and draft BL for OC 5RSG-90002. Please check the details and confirm.\n\nRegards,\nArjun")

    # 3 ── scanned, image-only documents → human review (+ Claude vision pre-read when enabled)
    f = OUT / "3_scanned_documents"; f.mkdir(exist_ok=True)
    scan_rows = [("Shipper", "APRIL FAR EAST (M) SDN BHD"), ("Consignee", "KTP CO., LTD"), ("Notify Party", NOTIFY),
                 ("Port of Loading", "PORT KLANG (WESTPORT), MALAYSIA"), ("Port of Discharge", "KARACHI, PAKISTAN"),
                 ("No. of Containers", "6 x 40'HC"), ("Gross Weight (KG)", "131,058 KG")]
    scanned_pdf(f / "SI_scan.pdf", "SHIPPING INSTRUCTION", scan_rows)
    scanned_pdf(f / "BL_scan.pdf", "BILL OF LADING (DRAFT)",
                [(k, "7 x 40'HC" if k == "No. of Containers" else v) for k, v in scan_rows])
    email(f / "email.txt", "REQUEST BL DRAFT _ PO 26999_ COATED IVORY BOARD__131MT",
          "Dear Team,\n\nAttached SI and draft BL for checking (scanned copies). Please advise.\n\nBest,\nFarah")

    # 4 ── wrong document in the BL slot
    f = OUT / "4_wrong_document"; f.mkdir(exist_ok=True)
    txt(f / "SI_order.txt", "SHIPPING INSTRUCTION", rows)
    (f / "BL_draft.txt").write_text(
        "PACKING LIST\n========================================\n\nShipper: APRIL FAR EAST (M) SDN BHD\n"
        "Consignee: KTP CO., LTD\nBooking Ref: SIN990003\n\nCarton No.   Net Wt (kg)   Gross Wt (kg)\nCTN-001      571           635\n",
        encoding="utf-8")
    email(f / "email.txt", "TO CONFIRM DOCS _ 5RSG-90004 _ BUSAN _ KTP CO., LTD",
          "Dear Team,\n\nPlease find attached the SI and draft BL. Kindly confirm the BL is in order.\n\nThanks")

    # 5 ── blank value → reviewer fills it in → report recomputes live
    f = OUT / "5_blank_value_human_fix"; f.mkdir(exist_ok=True)
    txt(f / "SI_blank.txt", "SHIPPING INSTRUCTION",
        [(k, "TBA" if k == "Gross Weight (KG)" else v) for k, v in rows])
    txt(f / "BL_blank.txt", "BILL OF LADING (DRAFT)", rows)
    email(f / "email.txt", "RE_ SI - SIN990005 - 5RSG-90005 - BUSAN_SOUTH KOREA",
          "Dear Team,\n\nPlease compare the SI and draft BL. Some SI fields were left blank by the customer; kindly confirm what we have.\n\nRegards")

    # 6 ── label the rules have never seen → Claude locates the field
    f = OUT / "6_unknown_label_ai"; f.mkdir(exist_ok=True)
    si6 = [(k, v) for k, v in rows if k != "Port of Loading"]
    si6.insert(3, ("Origin Terminal (Place of Receipt)", "SINGAPORE (SGSIN)"))
    txt(f / "SI_newformat.txt", "SHIPPING INSTRUCTION", si6)
    txt(f / "BL_newformat.txt", "BILL OF LADING (DRAFT)", rows)
    email(f / "email.txt", "TO CONFIRM DOCS _ 5RSG-90006 _ new customer template",
          "Dear Team,\n\nPlease check the draft BL against the SI — the customer is using a new SI template.\n\nThanks")

    # 7 ── no attachments, unusual wording → classification (Claude decides when rules are unsure)
    f = OUT / "7_ambiguous_email"; f.mkdir(exist_ok=True)
    email(f / "email.txt", "Quick one re Karachi",
          "Hey,\n\nBefore we lock anything in — what's the latest on the Karachi boxes? Customer keeps asking if the paperwork is final.\n\nCheers,\nDan")
    print("demo kit written to", OUT)


if __name__ == "__main__":
    main()
